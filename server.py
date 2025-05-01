from flask import Flask, request, jsonify, send_file, send_from_directory
from flask_cors import CORS
import os
from dotenv import load_dotenv
import openai
from pptx import Presentation
from pptx.util import Inches, Pt
import tempfile
import shutil

# Load environment variables
load_dotenv()

app = Flask(__name__, static_folder='sliding')
CORS(app)

# Configure OpenAI
openai.api_key = os.getenv('OPENAI_API_KEY')
if not openai.api_key:
    print("Warning: OPENAI_API_KEY not found in environment variables")

UPLOAD_FOLDER = 'uploads'
if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

@app.route('/')
def serve_index():
    return send_from_directory(app.static_folder, 'index.html')

@app.route('/<path:path>')
def serve_static(path):
    return send_from_directory(app.static_folder, path)

@app.route('/api/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No selected file'}), 400
    
    if file and file.filename.endswith(('.ppt', '.pptx')):
        filename = os.path.join(UPLOAD_FOLDER, file.filename)
        file.save(filename)
        return jsonify({'message': 'File uploaded successfully', 'filename': file.filename})
    
    return jsonify({'error': 'Invalid file type'}), 400

@app.route('/api/generate-content', methods=['POST'])
def generate_content():
    if not openai.api_key:
        return jsonify({'error': 'OpenAI API key not configured'}), 500

    data = request.json
    title = data.get('title')
    
    try:
        response = openai.ChatCompletion.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "You are a helpful presentation assistant. Generate a concise and engaging slide content based on the given title. Format the response as 'Title: [title]\nContent: [content]'"},
                {"role": "user", "content": f"Generate a title and content for a presentation slide about: {title}"}
            ],
            max_tokens=150,
            temperature=0.7
        )
        
        content = response.choices[0].message['content']
        return jsonify({'content': content})
    except Exception as e:
        print(f"OpenAI API error: {str(e)}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/save-presentation', methods=['POST'])
def save_presentation():
    data = request.json
    slides = data.get('slides', [])
    original_filename = data.get('filename')
    
    try:
        # Create a new presentation
        prs = Presentation()
        
        # Add slides
        for slide_data in slides:
            slide = prs.slides.add_slide(prs.slide_layouts[1])  # Using layout 1 for title and content
            
            # Add title
            title_shape = slide.shapes.title
            title_shape.text = slide_data.get('title', '')
            
            # Add content
            content_shape = slide.placeholders[1]
            content_shape.text = slide_data.get('content', '')
        
        # Save the presentation
        output_filename = 'presentation_edited.pptx'
        output_path = os.path.join(UPLOAD_FOLDER, output_filename)
        prs.save(output_path)
        
        return send_file(output_path, as_attachment=True)
    except Exception as e:
        print(f"Error saving presentation: {str(e)}")
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    if not openai.api_key:
        print("\nWarning: OpenAI API key not found!")
        print("Please add your API key to the .env file:")
        print("OPENAI_API_KEY=your_api_key_here\n")
    app.run(debug=True, port=5001) 